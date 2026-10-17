"""
CSV and Excel export of Bragg curves
"""

from typing import List

import pandas as pd

from .simulation import TrackResult
from .units import StoppingUnit

def _build_export_df(results: List[TrackResult], unit: StoppingUnit,
                     mass_thickness: bool, unit_in_name: bool = False) -> pd.DataFrame:
    """
    Shared table-building logic for CSV and Excel export
    """
    x_label = "x_g_per_cm2" if mass_thickness else "x_mm"
    suffix = f"_{unit.label}" if unit_in_name else ""

    labels: List[str] = []
    for r in results:
        label, k = r.name, 2
        while label in labels:
            label, k = f"{r.name}#{k}", k + 1
        labels.append(label)

    df, cols = None, []
    for r, label in zip(results, labels):
        e_col, s_col = f"E_{label}_MeV", f"dEdx_{label}{suffix}"
        cols += [e_col, s_col]
        d = pd.DataFrame({
            x_label: r.x_in(mass_thickness),
            e_col:   r.E,
            s_col:   r.dEdx_in(unit),
        })
        df = d if df is None else pd.merge(df, d, on=x_label, how="outer")

    df = df.sort_values(x_label).reset_index(drop=True)
    filled = df.set_index(x_label)[cols].interpolate(method="index", limit_area="inside")
    df[cols] = filled.to_numpy()
    return df


def export_csv(results: List[TrackResult], path: str,
               unit: StoppingUnit, mass_thickness: bool = False) -> None:
    _build_export_df(results, unit, mass_thickness).to_csv(path, index=False)


def _write_sheet(ws, df: pd.DataFrame, max_width: int, freeze_header: bool = False) -> None:
    """Write df to a worksheet with a styled header row and auto-sized columns"""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils.dataframe import dataframe_to_rows

    for row in dataframe_to_rows(df, index=False, header=True):
        ws.append(row)

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1E3A5F")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")

    for col in ws.columns:
        longest = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(longest + 2, max_width)

    if freeze_header:
        ws.freeze_panes = "A2"


_STOP_TEXT = {
    "stopped": "particle stopped",
    "model_limit": "electronic-stopping model limit",
    "x_stop": "depth limit (range NOT reached)",
    "window": "stopped in a window (never reaches the detector)",
}


def export_xlsx(results: List[TrackResult], path: str,
                unit: StoppingUnit, mass_thickness: bool = False,
                anodes=None) -> None:
    """
    Sheet "Bragg curves"
        One shared x-column, then alternating energy / dE/dx columns per
        particle — same structure as the CSV export so the two are
        interchangeable (the Excel headers additionally carry the unit)

    Sheet "Metadata"
        Entry/final energy, why the track ended, range, peak dE/dx and
        unit information for each particle, range straggling; with anodes
        (core.detector.AnodeLayout) also the energy deposited in each anode with its
        1-sigma width and the dE-E_res correlation
    """
    import openpyxl

    df = _build_export_df(results, unit, mass_thickness, unit_in_name=True)

    def meta(r: TrackResult) -> dict:
        row = {
            "Particle": r.name,
            "Beam energy (MeV)": float(f"{(r.E[0] if r.E_beam is None else r.E_beam):.6g}"),
            "Lost in windows (MeV)": float(f"{r.window_loss:.6g}"),
            "Entry energy (MeV)": float(f"{r.E[0]:.6g}"),
            "Final energy (MeV)": float(f"{r.E[-1]:.6g}"),
            "Track end": _STOP_TEXT[r.stop_reason],
            "Range (mm)": round(r.range_mm, 2),
            "Range straggling σ (mm)": (None if r.range_sigma_mm is None
                                        else float(f"{r.range_sigma_mm:.4g}")),
            # significant digits, not decimals: gas ranges are ~1e-4 g/cm²
            "Range (g/cm²)": float(f"{r.range_mass:.6g}"),
            f"Peak dE/dx ({unit.label})": round(r.peak_dEdx(unit), 5),
            "dE/dx unit": unit.label,
            "x-axis": "Mass thickness (g/cm²)" if mass_thickness else "Distance (mm)",
        }
        if anodes is not None:
            from .detector import anode_signals
            s = anode_signals(r, anodes)
            row.update({"Anode 1 dE (MeV)": float(f"{s.dE:.6g}"),
                        "Anode 1 dE σ (MeV)": None if s.sigma_dE is None else float(f"{s.sigma_dE:.4g}"),
                        "Anode 2 E_res (MeV)": float(f"{s.E_res:.6g}"),
                        "Anode 2 E_res σ (MeV)": None if s.sigma_E_res is None else float(f"{s.sigma_E_res:.4g}"),
                        "dE-E_res correlation": None if s.corr is None else round(s.corr, 4),
                        "Left after anode 2 (MeV)": float(f"{s.escaped:.6g}"),
                        "Anodes (gap / 1 / 2, mm)": f"{anodes.gap_mm:g} / {anodes.anode1_mm:g} / {anodes.anode2_mm:g}"})
        return row

    df_meta = pd.DataFrame([meta(r) for r in results])

    wb = openpyxl.Workbook()
    ws_data = wb.active
    ws_data.title = "Bragg curves"
    _write_sheet(ws_data, df, max_width=30, freeze_header=True)
    _write_sheet(wb.create_sheet("Metadata"), df_meta, max_width=35)
    wb.save(path)
